#!/usr/bin/env python3
"""pptx2md: convert PowerPoint decks to Markdown that reads well for LLMs.

Dropping Lecture5.pptx produces, next to it:

    Lecture5/
      Lecture5.md                     slide-by-slide text, tables, chart data, speaker notes
      images/
        slide3_img01.png              picture first seen on slide 3
        slide3,7,12_img01.png         same picture reused on slides 3, 7, 12 (saved once)
        slide3_full.png               full-slide snapshot, only for slides with a picture, chart,
                                      diagram, equation or drawn flowchart (text-only slides get none)

Every PNG gets a strip above the picture naming the slide(s) it belongs to.

Usage: pptx2md.py [--no-snapshots | --all-snapshots] [--no-open] [--no-notify] FILE...
"""

from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import io
import json
import logging
import math
import re
import shutil
import subprocess
import sys
import tempfile
import traceback
import uuid
import zipfile
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

from lxml import etree
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.shapes.shapetree import BaseShapeFactory, SlideShapeFactory, _NotesSlideShapeFactory

APP_NAME = "PPTX to Markdown"
LOG_PATH = Path.home() / "Library/Logs/pptx2md.log"
SNAPSHOT_WIDTH = 1920
MIN_IMAGE_PX = 10            # pictures smaller than this in both dimensions are spacer junk
MAX_SLIDE_LIST_CHARS = 180   # keeps "slide3,7,12_img01.png" under macOS's 255-byte filename limit
# A picture doesn't earn its slide a snapshot when it's an icon: tiny, or smallish and repeated on many slides.
ICON_MAX_AREA = 0.015
REPEATED_ICON_MAX_AREA = 0.10
REPEATED_ICON_MIN_SLIDES = 3
MIN_DRAWN_SHAPES = 3         # boxes/arrows/connectors on one slide that make it a drawn diagram

OOXML_SUFFIXES = {".pptx", ".pptm", ".ppsx", ".ppsm", ".potx", ".potm"}
LEGACY_SUFFIXES = {".ppt", ".pps", ".pot"}
PRESENTATION_MAIN_CT = "application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
    "dgm": "http://schemas.openxmlformats.org/drawingml/2006/diagram",
    "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
    "a14": "http://schemas.microsoft.com/office/drawing/2010/main",
    "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
    "p14": "http://schemas.microsoft.com/office/powerpoint/2010/main",
}


def q(tag: str) -> str:
    prefix, local = tag.split(":")
    return f"{{{NS[prefix]}}}{local}"


SHAPE_TAGS = {q("p:sp"), q("p:pic"), q("p:grpSp"), q("p:graphicFrame"), q("p:cxnSp")}
TITLE_TYPES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}
FURNITURE_TYPES = {PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.HEADER}
NOTES_SKIP_TYPES = FURNITURE_TYPES | {PP_PLACEHOLDER.SLIDE_IMAGE}
BULLET_BODY_TYPES = {PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT, PP_PLACEHOLDER.SUBTITLE, PP_PLACEHOLDER.VERTICAL_BODY}

log = logging.getLogger("pptx2md")


# --------------------------------------------------------------------------- small helpers

def norm_ws(text: str) -> str:
    return re.sub(r"\s+", " ", text.replace("\u00a0", " ")).strip()


def ph_type(shape):
    if not shape.is_placeholder:
        return None
    try:
        return shape.placeholder_format.type
    except Exception:
        return None


def geom(shape):
    def val(name):
        try:
            return getattr(shape, name) or 0
        except Exception:
            return 0
    return val("left"), val("top"), val("width"), val("height")


def iter_shapes(container_elm, parent, factory):
    """Yield shapes in a spTree/grpSp, including ones python-pptx skips inside mc:AlternateContent
    (PowerPoint wraps text boxes containing equations that way)."""
    for child in container_elm.iterchildren():
        if child.tag in SHAPE_TAGS:
            yield factory(child, parent)
        elif child.tag == q("mc:AlternateContent"):
            branch = child.find(q("mc:Choice"))
            if branch is None:
                branch = child.find(q("mc:Fallback"))
            for sub in branch.iterchildren() if branch is not None else ():
                if sub.tag in SHAPE_TAGS:
                    yield factory(sub, parent)


def group_area_scale(grp_elm) -> float:
    """Group children are positioned in the group's own child coordinates; this converts their areas to slide units."""
    xfrm = grp_elm.find(q("p:grpSpPr") + "/" + q("a:xfrm"))
    try:
        ext, chext = xfrm.find(q("a:ext")), xfrm.find(q("a:chExt"))
        return (int(ext.get("cx")) * int(ext.get("cy"))) / (int(chext.get("cx")) * int(chext.get("cy")))
    except (AttributeError, TypeError, ValueError, ZeroDivisionError):
        return 1.0


def is_drawn_shape(sp_elm) -> bool:
    """A visibly drawn autoshape (box, arrow, line, flowchart symbol), as opposed to a text box or placeholder."""
    nv = sp_elm.find(q("p:nvSpPr"))
    if nv is None or nv.find(q("p:nvPr") + "/" + q("p:ph")) is not None:
        return False
    cnv = nv.find(q("p:cNvSpPr"))
    if cnv is not None and cnv.get("txBox") in ("1", "true"):
        return False
    sppr = sp_elm.find(q("p:spPr"))
    if sppr is None:
        return False
    fills = {q("a:solidFill"), q("a:gradFill"), q("a:blipFill"), q("a:pattFill")}
    styled = sp_elm.find(q("p:style")) is not None  # theme style draws fill/outline unless overridden
    fill_visible = any(c.tag in fills for c in sppr) or (styled and sppr.find(q("a:noFill")) is None)
    ln = sppr.find(q("a:ln"))
    if ln is None:
        line_visible = styled
    else:
        line_visible = any(c.tag in fills for c in ln) or (styled and ln.find(q("a:noFill")) is None)
    return fill_visible or line_visible


def reading_order(shapes, slide_h):
    """Titles first, then rows top to bottom; shapes whose tops are within 5% of the slide height of
    the row's first shape share a row and are read left to right (handles two-column layouts)."""
    tol = slide_h * 0.05
    titles = [s for s in shapes if ph_type(s) in TITLE_TYPES]
    rest = sorted((s for s in shapes if ph_type(s) not in TITLE_TYPES), key=lambda s: (geom(s)[1], geom(s)[0]))
    rows, row_top = [], None
    for s in rest:
        top = geom(s)[1]
        if row_top is None or top > row_top + tol:
            rows.append([])
            row_top = top
        rows[-1].append(s)
    return titles + [s for row in rows for s in sorted(row, key=lambda s: geom(s)[0])]


def fmt_number(v: str) -> str:
    try:
        x = float(v)
    except (TypeError, ValueError):
        return v or ""
    if math.isfinite(x) and x == int(x) and abs(x) < 1e15:
        return str(int(x))
    return f"{x:.10g}"


def md_cell(text: str) -> str:
    return text.replace("|", "\\|").replace("\n", "<br>").strip()


def md_table(rows: list[list[str]]) -> str:
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    out = ["| " + " | ".join(md_cell(c) for c in rows[0]) + " |", "|" + "---|" * width]
    out += ["| " + " | ".join(md_cell(c) for c in r) + " |" for r in rows[1:]]
    return "\n".join(out)


def slide_list_str(nums: list[int]) -> str:
    s = ",".join(map(str, nums))
    if len(s) <= MAX_SLIDE_LIST_CHARS:
        return s
    runs, start, prev = [], nums[0], nums[0]
    for n in nums[1:] + [None]:
        if n is not None and n == prev + 1:
            prev = n
            continue
        runs.append(str(start) if start == prev else f"{start}-{prev}")
        if n is not None:
            start = prev = n
    s = ",".join(runs)
    return s if len(s) <= MAX_SLIDE_LIST_CHARS else f"{nums[0]}-{nums[-1]}_{len(nums)}slides"


def slides_text(nums: list[int]) -> str:
    """'3, 7, 12' for short lists; long lists collapse consecutive runs: '1–40, 42, 44–84'."""
    if len(nums) <= 10:
        return ", ".join(map(str, nums))
    runs, start, prev = [], nums[0], nums[0]
    for n in nums[1:] + [None]:
        if n is not None and n == prev + 1:
            prev = n
            continue
        runs.append(str(start) if start == prev else f"{start}–{prev}")
        if n is not None:
            start = prev = n
    return ", ".join(runs)


def slides_label(nums: list[int]) -> str:
    return f"Slide {nums[0]}" if len(nums) == 1 else "Slides " + slides_text(nums)


def clean_alt(text: str | None) -> str:
    if not text:
        return ""
    text = re.sub(r"\s*Description automatically generated.*$", "", text, flags=re.S | re.I)
    text = norm_ws(text)
    if re.match(r"^Google Shape;", text) or re.fullmatch(r"[\w\-\s().,]+\.(png|jpe?g|gif|bmp|tiff?|emf|wmf|svg|webp)", text, re.I):
        return ""
    if re.fullmatch(r"(Picture|Image|Graphic|Content Placeholder|Picture Placeholder)\s*\d*", text, re.I):
        return ""
    return text


# --------------------------------------------------------------------------- text rendering

def _bullet_from_ppr(ppr):
    if ppr is None:
        return None
    for child in ppr:
        if child.tag == q("a:buNone"):
            return "none"
        if child.tag in (q("a:buChar"), q("a:buBlip")):
            return "bullet"
        if child.tag == q("a:buAutoNum"):
            return "number"
    return None


def bullet_kind(p_el, shape, body_style) -> str | None:
    """Resolve whether a paragraph is bulleted, following PowerPoint's inheritance chain:
    paragraph -> shape list style -> layout placeholder -> master placeholder -> master body style."""
    ppr = p_el.find(q("a:pPr"))
    lvl = int(ppr.get("lvl", "0")) if ppr is not None else 0
    kind = _bullet_from_ppr(ppr)
    if kind or shape is None:
        return None if kind in (None, "none") else kind

    lvl_tag = q(f"a:lvl{lvl + 1}pPr")
    node, hops = shape, 0
    while node is not None and hops < 4:
        txbody = node._element.find(q("p:txBody"))
        lst = txbody.find(q("a:lstStyle")) if txbody is not None else None
        kind = _bullet_from_ppr(lst.find(lvl_tag)) if lst is not None else None
        if kind:
            return None if kind == "none" else kind
        if not node.is_placeholder:
            return None  # plain text boxes inherit no bullets
        try:
            node = node._base_placeholder
        except Exception:
            node = None
        hops += 1

    if ph_type(shape) in BULLET_BODY_TYPES and body_style is not None:
        kind = _bullet_from_ppr(body_style.find(lvl_tag))
        return None if kind in (None, "none") else kind
    return None


def _run_format(rpr, part):
    if rpr is None:
        return (False, False, 0, None)
    bold = rpr.get("b") in ("1", "true")
    ital = rpr.get("i") in ("1", "true")
    try:
        base = int(rpr.get("baseline", "0"))
    except ValueError:
        base = 0
    href = None
    link = rpr.find(q("a:hlinkClick"))
    if link is not None and link.get(q("r:id")) and part is not None:
        try:
            rel = part.rels[link.get(q("r:id"))]
            if rel.is_external:
                href = rel.target_ref
        except KeyError:
            pass
    return (bold, ital, (base > 0) - (base < 0), href)


def _math_text(el) -> str:
    return "".join(t.text or "" for t in el.iter(q("m:t")))


def _paragraph_segments(p_el, part):
    """Yield (text, format) for every piece of a paragraph: runs, fields, line breaks, equations."""
    for child in p_el.iterchildren():
        tag = child.tag
        if tag in (q("a:r"), q("a:fld")):
            t = child.find(q("a:t"))
            if t is not None and t.text:
                yield t.text.replace("\x0b", "\n"), _run_format(child.find(q("a:rPr")), part)
        elif tag == q("a:br"):
            yield "\n", None
        elif tag == q("a14:m"):
            yield ("math", _math_text(child)), None
        elif tag == q("mc:AlternateContent"):
            choice = child.find(q("mc:Choice"))
            for sub in choice.iterchildren() if choice is not None else ():
                if sub.tag == q("a14:m"):
                    yield ("math", _math_text(sub)), None
                elif sub.tag in (q("a:r"), q("a:fld")):
                    t = sub.find(q("a:t"))
                    if t is not None and t.text:
                        yield t.text, _run_format(sub.find(q("a:rPr")), part)


def _wrap(text, fmt):
    if fmt is None or not text.strip():
        return text
    bold, ital, vert, href = fmt
    lead = text[: len(text) - len(text.lstrip())]
    trail = text[len(text.rstrip()):]
    core = text.strip()
    if vert > 0:
        core = f"<sup>{core}</sup>"
    elif vert < 0:
        core = f"<sub>{core}</sub>"
    if bold and ital:
        core = f"***{core}***"
    elif bold:
        core = f"**{core}**"
    elif ital:
        core = f"*{core}*"
    if href:
        core = f"[{core}]({href})"
    return lead + core + trail


def render_paragraph(p_el, part):
    """Return (markdown, plain_text) for one a:p."""
    merged: list[list] = []  # [text, fmt]
    for text, fmt in _paragraph_segments(p_el, part):
        if isinstance(text, tuple):  # equation
            math_txt = text[1].strip()
            if math_txt:
                merged.append([f"`{math_txt}`", "raw", math_txt])
            continue
        for i, piece in enumerate(text.split("\n")):
            if i:
                merged.append(["\n", None])
            if piece:
                if merged and merged[-1][1] == fmt and fmt != "raw":
                    merged[-1][0] += piece
                else:
                    merged.append([piece, fmt])
    md, plain = [], []
    for item in merged:
        if item[1] == "raw":
            md.append(item[0])
            plain.append(item[2])
        else:
            md.append(_wrap(item[0], item[1]))
            plain.append(item[0])
    return "".join(md), "".join(plain)


def render_text_body(txbody, part, shape=None, body_style=None):
    """Render a txBody to markdown, keeping bullet levels. Returns (markdown, plain_text)."""
    items = []  # (is_list, markdown)
    plains = []
    for p_el in txbody.iterfind(q("a:p")):
        md, plain = render_paragraph(p_el, part)
        if not plain.strip():
            items.append((False, None))
            continue
        plains.append(plain)
        ppr = p_el.find(q("a:pPr"))
        lvl = int(ppr.get("lvl", "0")) if ppr is not None else 0
        kind = bullet_kind(p_el, shape, body_style)
        md = md.strip()
        if kind:
            indent = "  " * lvl
            marker = "1." if kind == "number" else "-"
            md = indent + marker + " " + md.replace("\n", "\n" + indent + " " * (len(marker) + 1))
            items.append((True, md))
        else:
            if md.startswith(("#", "---", "===")):
                md = "\\" + md
            items.append((False, md))
    out, prev_list = [], None
    for is_list, md in items:
        if md is None:
            prev_list = None
            continue
        if out:
            out.append("\n" if (is_list and prev_list) else "\n\n")
        out.append(md)
        prev_list = is_list
    return "".join(out), "\n".join(plains)


def shape_text(shape, part, body_style):
    txbody = shape._element.find(q("p:txBody"))
    if txbody is None:
        return "", ""
    return render_text_body(txbody, part, shape, body_style)


def title_text(shape) -> str:
    txbody = shape._element.find(q("p:txBody"))
    if txbody is None:
        return ""
    return norm_ws(" ".join(render_paragraph(p, None)[1] for p in txbody.iterfind(q("a:p"))))


# --------------------------------------------------------------------------- tables, charts, SmartArt

def table_md(shape, part) -> str:
    rows = []
    for row in shape.table.rows:
        cells = []
        for cell in row.cells:
            if cell.is_spanned:
                cells.append("")
                continue
            txbody = cell._tc.find(q("a:txBody"))
            cells.append(render_text_body(txbody, part)[0].replace("\n\n", "\n") if txbody is not None else "")
        rows.append(cells)
    return md_table(rows)


def _pt_map(el):
    """{idx: value} from a chart c:cat / c:val / c:xVal / c:yVal element."""
    if el is None:
        return {}
    lvl = el.find(".//" + q("c:lvl"))  # multi-level categories: take the innermost level
    scope = lvl if lvl is not None else el
    out = {}
    for pt in scope.iter(q("c:pt")):
        v = pt.find(q("c:v"))
        out[int(pt.get("idx", len(out)))] = v.text if v is not None else ""
    return out


def chart_md(shape) -> str:
    cs = shape.chart._chartSpace
    chart = cs.find(q("c:chart"))
    title_el = chart.find(q("c:title")) if chart is not None else None
    title = norm_ws(" ".join(t.text or "" for t in title_el.iter(q("a:t")))) if title_el is not None else ""
    plot_area = chart.find(q("c:plotArea")) if chart is not None else None
    parts = []
    for plot in plot_area.iterchildren() if plot_area is not None else ():
        local = etree.QName(plot).localname
        if not local.endswith("Chart"):
            continue
        kind = re.sub(r"(3D)?Chart$", "", local)
        series = []
        for i, ser in enumerate(plot.iterfind(q("c:ser")), 1):
            tx = ser.find(q("c:tx"))
            name = norm_ws(" ".join(v.text or "" for v in tx.iter(q("c:v")))) if tx is not None else ""
            if not name and tx is not None:
                name = norm_ws(" ".join(t.text or "" for t in tx.iter(q("a:t"))))
            cats = _pt_map(ser.find(q("c:cat")) if ser.find(q("c:cat")) is not None else ser.find(q("c:xVal")))
            vals = _pt_map(ser.find(q("c:val")) if ser.find(q("c:val")) is not None else ser.find(q("c:yVal")))
            series.append((name or f"Series {i}", cats, vals))
        if not series:
            continue
        if kind in ("scatter", "bubble"):
            for name, xs, ys in series:
                idxs = sorted(set(xs) | set(ys))
                rows = [["X", f"{name} (Y)"]] + [[fmt_number(xs.get(i, str(i + 1))), fmt_number(ys.get(i, ""))] for i in idxs]
                parts.append(f"*{kind} chart — {name}*\n\n" + md_table(rows))
        else:
            cats = series[0][1]
            idxs = sorted(set(cats).union(*(set(v) for _, _, v in series)))
            rows = [["Category"] + [s[0] for s in series]]
            rows += [[cats.get(i, str(i + 1))] + [fmt_number(v.get(i, "")) for _, _, v in series] for i in idxs]
            parts.append(f"*{kind} chart*\n\n" + md_table(rows))
    head = f"**Chart: {title}**" if title else "**Chart**"
    return head + "\n\n" + ("\n\n".join(parts) if parts else "_(no readable chart data)_")


def smartart_md(shape, part) -> str:
    rel_ids = shape._element.find(".//" + q("dgm:relIds"))
    if rel_ids is None:
        return ""
    data = etree.fromstring(part.related_part(rel_ids.get(q("r:dm"))).blob)
    texts, parent, order = {}, {}, {}
    for pt in data.iter(q("dgm:pt")):
        if pt.get("type", "node") not in ("node", "asst"):
            continue
        t = pt.find(q("dgm:t"))
        if t is None:
            continue
        paras = [norm_ws("".join(x.text or "" for x in p.iter(q("a:t")))) for p in t.iter(q("a:p"))]
        text = " ".join(s for s in paras if s)
        if text:
            texts[pt.get("modelId")] = text
    for cx in data.iter(q("dgm:cxn")):
        if cx.get("type", "parOf") == "parOf":
            parent[cx.get("destId")] = cx.get("srcId")
            order[cx.get("destId")] = int(cx.get("srcOrd", "0"))

    def depth(mid):
        d, seen = 0, set()
        while mid in parent and mid not in seen:
            seen.add(mid)
            mid = parent[mid]
            if mid in texts:
                d += 1
        return d

    def path_key(mid):
        chain, seen = [], set()
        while mid is not None and mid not in seen:
            seen.add(mid)
            chain.append(order.get(mid, 0))
            mid = parent.get(mid)
        return chain[::-1]

    lines = ["  " * depth(m) + "- " + texts[m] for m in sorted(texts, key=path_key)]
    return "**Diagram (SmartArt):**\n\n" + "\n".join(lines) if lines else ""


# --------------------------------------------------------------------------- extraction

@dataclass
class ImageInfo:
    key: str
    blob: bytes
    ext: str
    slides: list = field(default_factory=list)
    pil: Image.Image | None = None
    size: tuple | None = None
    ordinal: int = 0
    name: str = ""
    raw_only: bool = False


@dataclass
class SlideData:
    number: int
    title: str
    hidden: bool
    section: str | None
    blocks: list            # ("md", str) | ("img", key, alt)
    offslide: list
    notes: list             # same block format
    notes_plain: str
    visual: list = field(default_factory=list)            # ("picture", key, area) | ("shape",) | ("chart",) ...
    snapshot_reasons: list = field(default_factory=list)

    def image_keys(self):
        for blk in self.blocks + self.offslide + self.notes:
            if blk[0] == "img":
                yield blk[1]


class Extractor:
    def __init__(self, prs):
        self.prs = prs
        self.W, self.H = prs.slide_width or 12192000, prs.slide_height or 6858000
        self.images: dict[str, ImageInfo] = {}
        self.footers: set[str] = set()

    # -- image references
    def image_blocks(self, shape_elm, part, alt):
        blocks = []
        for blip in shape_elm.iter(q("a:blip")):
            if blip.getparent() is not None and blip.getparent().tag == q("a:buBlip"):
                continue  # picture bullet
            rid = blip.get(q("r:embed"))
            if not rid:
                link = blip.get(q("r:link"))
                if link and link in part.rels:
                    blocks.append(("md", f"_(Linked image not embedded in the file: {part.rels[link].target_ref})_"))
                continue
            try:
                img_part = part.related_part(rid)
                blob = img_part.blob
            except Exception:
                continue
            key = hashlib.sha1(blob).hexdigest()
            if key not in self.images:
                self.images[key] = ImageInfo(key, blob, img_part.partname.ext.lower())
            blocks.append(("img", key, alt))
        return blocks

    # -- one shape (recursive for groups)
    def shape_blocks(self, shape, part, body_style, visual=None, area_scale=1.0):
        """Markdown/image blocks for a shape. When `visual` is a list, also record the non-text content
        the shape carries (pictures, charts, drawn shapes…), which decides whether the slide gets a snapshot."""
        if ph_type(shape) in FURNITURE_TYPES:
            return []
        elm = shape._element
        if elm.tag == q("p:grpSp"):
            scale = area_scale * group_area_scale(elm)
            children = list(iter_shapes(elm, shape.shapes, BaseShapeFactory))
            out = []
            for child in reading_order(children, self.H):
                out += self.shape_blocks(child, part, body_style, visual, scale)
            return out

        out = []
        cnv = elm.find(".//" + q("p:cNvPr"))
        alt = clean_alt((cnv.get("descr") or cnv.get("title")) if cnv is not None else "")
        if elm.tag != q("p:graphicFrame") or not getattr(shape, "has_chart", False):
            imgs = self.image_blocks(elm, part, alt)
            out += imgs
            if visual is not None:
                _, _, w, h = geom(shape)
                area = w * h * area_scale / (self.W * self.H)
                visual += [("picture", b[1], area) for b in imgs if b[0] == "img"]

        if elm.tag == q("p:sp"):
            md, plain = shape_text(shape, part, body_style)
            if plain.strip():
                out.append(("md", md))
            if visual is not None and is_drawn_shape(elm):
                visual.append(("shape",))
        elif elm.tag == q("p:cxnSp"):
            if visual is not None:
                visual.append(("shape",))
        elif elm.tag == q("p:graphicFrame"):
            uri = elm.find(".//" + q("a:graphicData"))
            uri = uri.get("uri", "") if uri is not None else ""
            try:
                if shape.has_table:
                    out.append(("md", table_md(shape, part)))
                elif shape.has_chart:
                    out.append(("md", chart_md(shape)))
                    if visual is not None:
                        visual.append(("chart",))
                elif uri.endswith("/diagram"):
                    md = smartart_md(shape, part)
                    if md:
                        out.append(("md", md))
                    if visual is not None:
                        visual.append(("diagram",))
                elif visual is not None and elm.find(".//" + q("p:oleObj")) is not None:
                    visual.append(("embedded object",))
            except Exception as e:
                log.warning("could not read graphic frame: %s", e)
                out.append(("md", "_(This slide has a table/chart/diagram that could not be read — see the full-slide snapshot.)_"))
        return out

    def is_offslide(self, shape):
        left, top, w, h = geom(shape)
        return left >= self.W or top >= self.H or left + w <= 0 or top + h <= 0

    def top_shapes(self, slide):
        return list(iter_shapes(slide.shapes._spTree, slide.shapes, SlideShapeFactory))

    def find_footers(self):
        """Text boxes repeated verbatim near the bottom of many slides (e.g. 'Lecture 10: ...')."""
        counts = Counter()
        for slide in self.prs.slides:
            seen = set()
            for sh in self.top_shapes(slide):
                if not sh.is_placeholder and sh.has_text_frame:
                    t = norm_ws(sh.text_frame.text)
                    if t:
                        seen.add(t)
            counts.update(seen)
        n = len(self.prs.slides)
        threshold = max(3, math.ceil(0.25 * n))
        self.footers = {t for t, c in counts.items() if c >= threshold and not t.isdigit()}

    def is_footer(self, shape, slide_no):
        if shape.is_placeholder or not shape.has_text_frame:
            return False
        _, top, _, _ = geom(shape)
        if not (top >= 0.85 * self.H or self.is_offslide(shape)):
            return False
        t = norm_ws(shape.text_frame.text)
        return t in self.footers or (t.isdigit() and int(t) == slide_no)

    def sections(self):
        lst = self.prs.part._element.find(".//" + q("p14:sectionLst"))
        out = {}
        if lst is not None:
            for sec in lst.iterfind(q("p14:section")):
                for sid in sec.iter(q("p14:sldId")):
                    out[int(sid.get("id"))] = sec.get("name", "")
        return out

    def notes_blocks(self, slide):
        if not slide.has_notes_slide:
            return [], ""
        ns = slide.notes_slide
        style = None
        try:
            style = ns.part.notes_master._element.find(q("p:notesStyle"))
        except Exception:
            pass
        blocks, plains, raw = [], [], []
        shapes = list(iter_shapes(ns.shapes._spTree, ns.shapes, _NotesSlideShapeFactory))
        for sh in shapes:
            if ph_type(sh) in NOTES_SKIP_TYPES:
                continue
            if sh._element.tag == q("p:sp"):
                md, plain = shape_text(sh, ns.part, style)
                if plain.strip():
                    blocks.append(("md", md))
                    plains.append(plain)
                    raw.append(sh.text_frame.text)
            if sh._element.tag in (q("p:pic"), q("p:sp")):
                blocks += self.image_blocks(sh._element, ns.part, "")
            elif sh._element.tag == q("p:graphicFrame") and getattr(sh, "has_table", False):
                blocks.append(("md", table_md(sh, ns.part)))

        # Speaker notes must never be lost: every line python-pptx sees must appear in our output.
        got = norm_ws("\n".join(plains))
        missing = [ln for t in raw for ln in re.split(r"[\n\x0b]", t) if norm_ws(ln) and norm_ws(ln) not in got]
        if missing:
            log.warning("slide %s: notes formatting fallback for %d line(s)", slide.slide_id, len(missing))
            blocks.append(("md", "**Raw notes text:**\n\n" + "\n".join(raw)))
            plains.append("\n".join(raw))
        return blocks, "\n".join(plains)

    def extract(self) -> list[SlideData]:
        self.find_footers()
        sections = self.sections()
        master_cache = {}
        out = []
        for no, slide in enumerate(self.prs.slides, 1):
            master = slide.slide_layout.slide_master
            body_style = master_cache.setdefault(id(master), master._element.find(q("p:txStyles") + "/" + q("p:bodyStyle")))
            shapes = reading_order(self.top_shapes(slide), self.H)
            title = ""
            blocks, offslide, visual = [], [], []

            bg_fill = slide._element.find(q("p:cSld") + "/" + q("p:bg"))
            if bg_fill is not None:
                blocks += [(b[0], b[1], "background image") if b[0] == "img" else b
                           for b in self.image_blocks(bg_fill, slide.part, "")]

            for sh in shapes:
                if not title and ph_type(sh) in TITLE_TYPES:
                    title = title_text(sh)
                    blocks += [b for b in self.image_blocks(sh._element, slide.part, "")]
                    continue
                if self.is_footer(sh, no):
                    continue
                if self.is_offslide(sh):  # not visible, so it never earns a snapshot
                    offslide += self.shape_blocks(sh, slide.part, body_style)
                else:
                    blocks += self.shape_blocks(sh, slide.part, body_style, visual)
            if next(slide._element.iter(q("a14:m")), None) is not None:
                visual.append(("equation",))
            if next(slide._element.iter(q("p:contentPart")), None) is not None:
                visual.append(("ink drawing",))

            notes, notes_plain = self.notes_blocks(slide)
            out.append(SlideData(no, title, slide._element.get("show") == "0",
                                 sections.get(slide.slide_id), blocks, offslide, notes, notes_plain, visual))
        return out


def snapshot_reasons(sd: SlideData, images: dict) -> list[str]:
    """Why this slide deserves a full-slide snapshot; empty for text-only slides (tables count as text)."""
    reasons, drawn = set(), 0
    for item in sd.visual:
        if item[0] == "picture":
            _, key, area = item
            n_slides = len(images[key].slides)
            if not n_slides:
                continue  # spacer image that was dropped
            icon = area < ICON_MAX_AREA or (n_slides >= REPEATED_ICON_MIN_SLIDES and area < REPEATED_ICON_MAX_AREA)
            if not icon:
                reasons.add("picture")
        elif item[0] == "shape":
            drawn += 1
        else:
            reasons.add(item[0])
    if drawn >= MIN_DRAWN_SHAPES:
        reasons.add("drawn diagram")
    return sorted(reasons)


# --------------------------------------------------------------------------- images

def load_font(size):
    for path, idx in (("/System/Library/Fonts/Helvetica.ttc", 1), ("/System/Library/Fonts/SFNS.ttf", 0),
                      ("/Library/Fonts/Arial.ttf", 0)):
        try:
            return ImageFont.truetype(path, size, index=idx)
        except Exception:
            continue
    return ImageFont.load_default(size=size)


def decode_image(info: ImageInfo, workdir: Path):
    im = None
    try:
        im = Image.open(io.BytesIO(info.blob))
        im.seek(0)
        im.load()
    except Exception:
        # EMF/WMF/WDP and friends: let macOS ImageIO try
        src = workdir / f"{info.key}{info.ext and '.' + info.ext}"
        dst = workdir / f"{info.key}.png"
        src.write_bytes(info.blob)
        subprocess.run(["/usr/bin/sips", "-s", "format", "png", str(src), "--out", str(dst)],
                       capture_output=True, timeout=120)
        if dst.exists() and dst.stat().st_size:
            try:
                im = Image.open(dst)
                im.load()
            except Exception:
                im = None
    if im is None:
        info.raw_only = True
        return
    try:
        has_alpha = im.mode in ("RGBA", "LA", "PA") or (im.mode == "P" and "transparency" in im.info)
        im = im.convert("RGBA" if has_alpha else "RGB")
    except Exception:
        im = im.convert("RGBA")
    info.pil, info.size = im, im.size


def add_label(im: Image.Image, text: str) -> Image.Image:
    w, h = im.size
    fs = int(max(14, min(48, w * 0.025)))
    font = load_font(fs)
    pad = max(6, int(fs * 0.45))
    probe = ImageDraw.Draw(Image.new("RGB", (1, 1)))
    max_w = max(w - 2 * pad, fs * 6)
    lines, cur = [], ""
    for word in text.split(" "):
        cand = f"{cur} {word}".strip()
        if cur and probe.textlength(cand, font=font) > max_w:
            lines.append(cur)
            cur = word
        else:
            cur = cand
    lines.append(cur)
    ascent, descent = font.getmetrics()
    line_h = ascent + descent
    strip_h = len(lines) * line_h + 2 * pad
    canvas_w = max(w, int(max(probe.textlength(l, font=font) for l in lines)) + 2 * pad)
    rgba = im.mode == "RGBA"
    canvas = Image.new(im.mode, (canvas_w, h + strip_h), (255, 255, 255, 0) if rgba else (255, 255, 255))
    draw = ImageDraw.Draw(canvas)
    draw.rectangle([0, 0, canvas_w, strip_h], fill=(31, 41, 55, 255) if rgba else (31, 41, 55))
    for i, line in enumerate(lines):
        draw.text((pad, pad + i * line_h), line, font=font, fill=(255, 255, 255, 255) if rgba else (255, 255, 255))
    canvas.paste(im, (0, strip_h))
    return canvas


# --------------------------------------------------------------------------- PowerPoint automation

class PowerPointError(RuntimeError):
    pass


def _as(s) -> str:
    return '"' + str(s).replace("\\", "\\\\").replace('"', '\\"') + '"'


def powerpoint_save(src: Path, dst: Path, save_as: str, timeout: int = 300):
    """Open src in PowerPoint, save it as dst using the given 'save as ...' format, close it.

    PowerPoint is sandboxed: it can only write to files AppleScript hands it an extension for,
    which only works for files that already exist, so dst is created empty first."""
    if not Path("/Applications/Microsoft PowerPoint.app").exists():
        raise PowerPointError("Microsoft PowerPoint is not installed")
    dst.write_bytes(b"")
    script = f"""
set srcPath to {_as(src)}
set dstPath to {_as(dst)}
set wasRunning to application "Microsoft PowerPoint" is running
with timeout of {timeout} seconds
  tell application "Microsoft PowerPoint"
    open (POSIX file srcPath)
    set pres to missing value
    repeat 240 times
      repeat with i from 1 to (count of presentations)
        if (full name of presentation i) ends with {_as(src.name)} then set pres to presentation i
      end repeat
      if pres is not missing value then exit repeat
      delay 0.25
    end repeat
    if pres is missing value then error "PowerPoint did not open the deck"
    save pres in (POSIX file dstPath) as {save_as}
    close pres saving no
    if not wasRunning then
      if (count of presentations) is 0 then quit
    end if
  end tell
end timeout
"""
    try:
        res = subprocess.run(["/usr/bin/osascript", "-"], input=script, text=True,
                             capture_output=True, timeout=timeout + 60)
    except subprocess.TimeoutExpired:
        raise PowerPointError("PowerPoint took too long (is a dialog open in PowerPoint?)")
    if res.returncode != 0:
        err = res.stderr.strip()
        if "-1743" in err or "Not authorized" in err:
            raise PowerPointError(f"{APP_NAME} is not allowed to control PowerPoint. Enable it in System Settings → "
                                  "Privacy & Security → Automation.")
        raise PowerPointError(f"PowerPoint error: {err}")
    if not dst.exists() or dst.stat().st_size == 0:
        raise PowerPointError("PowerPoint did not write the exported file")


def render_snapshots(pptx_path: Path, n_slides: int, has_hidden: bool, workdir: Path, wanted: set[int]):
    """Yield (slide_number, image) for the wanted slides."""
    import pypdfium2 as pdfium

    token = uuid.uuid4().hex[:8]
    deck = workdir / f"pptx2md-{token}-snap.pptx"
    if has_hidden:
        prs = Presentation(pptx_path)  # PDF export skips hidden slides, so unhide them in a copy
        for slide in prs.slides:
            slide._element.attrib.pop("show", None)
        prs.save(deck)
    else:
        shutil.copyfile(pptx_path, deck)
    pdf_path = workdir / f"pptx2md-{token}-snap.pdf"
    powerpoint_save(deck, pdf_path, "save as PDF")
    pdf = pdfium.PdfDocument(str(pdf_path))
    try:
        if len(pdf) != n_slides:
            raise PowerPointError(f"PowerPoint exported {len(pdf)} pages for {n_slides} slides")
        for no in sorted(wanted):
            page = pdf[no - 1]
            w, _ = page.get_size()
            img = page.render(scale=SNAPSHOT_WIDTH / w).to_pil().convert("RGB")
            page.close()
            yield no, img
    finally:
        pdf.close()


# --------------------------------------------------------------------------- input preparation

def prepare_input(src: Path, workdir: Path) -> Path:
    """Return a path python-pptx can open."""
    suffix = src.suffix.lower()
    token = uuid.uuid4().hex[:8]
    if suffix == ".pptx":
        return src
    if suffix in OOXML_SUFFIXES:
        # .pptm/.ppsx/.potx are the same format with a different main content type
        out = workdir / f"pptx2md-{token}.pptx"
        with zipfile.ZipFile(src) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename == "[Content_Types].xml":
                    data = re.sub(rb'(PartName="/ppt/presentation\.xml"\s+ContentType=")[^"]+"',
                                  rb'\g<1>' + PRESENTATION_MAIN_CT.encode() + b'"', data)
                zout.writestr(item, data)
        return out
    if suffix in LEGACY_SUFFIXES:
        tmp_src = workdir / f"pptx2md-{token}.ppt"  # .pps would open straight into a slide show
        shutil.copyfile(src, tmp_src)
        out = workdir / f"pptx2md-{token}-converted.pptx"
        powerpoint_save(tmp_src, out, "save as Open XML presentation")
        return out
    raise ValueError(f"{src.name} is not a PowerPoint file")


def make_outdir(src: Path) -> Path:
    base = src.parent
    probe = base / f".pptx2md-{uuid.uuid4().hex[:6]}"
    try:
        probe.mkdir()
        probe.rmdir()
    except OSError:
        base = Path.home() / "Desktop" / "PPTX Conversions"
        base.mkdir(parents=True, exist_ok=True)
    candidate, n = base / src.stem, 2
    while candidate.exists():
        candidate = base / f"{src.stem} ({n})"
        n += 1
    (candidate / "images").mkdir(parents=True)
    return candidate


# --------------------------------------------------------------------------- markdown assembly

def build_markdown(src: Path, prs, slides: list[SlideData], images: dict[str, ImageInfo],
                   snapped: set[int], footers: set, warnings: list) -> str:
    title = (prs.core_properties.title or "").strip()
    if not title or title.lower() in ("powerpoint presentation", "presentation"):
        title = src.stem
    hidden = [s.number for s in slides if s.hidden]
    with_notes = [s.number for s in slides if s.notes_plain.strip()]
    used = [i for i in images.values() if i.slides]
    fm = [
        "---",
        f"source: {json.dumps(src.name, ensure_ascii=False)}",
        f"converted: {json.dumps(dt.datetime.now().strftime('%Y-%m-%d %H:%M'))}",
        f"slides: {len(slides)}",
        f"hidden_slides: {json.dumps(hidden)}",
        f"slides_with_speaker_notes: {len(with_notes)}",
        f"unique_images: {len(used)}",
        f"full_slide_snapshots: {json.dumps(sorted(snapped))}",
    ]
    if footers:
        fm.append(f"omitted_repeating_footer: {json.dumps(sorted(footers), ensure_ascii=False)}")
    if warnings:
        fm.append(f"warnings: {json.dumps(warnings, ensure_ascii=False)}")
    fm.append("---")

    out = ["\n".join(fm), f"# {title}",
           "> Converted from PowerPoint. Each `## Slide N` section holds that slide's text in reading order, "
           "its images (each PNG is labeled with its slide number), and the presenter's **Speaker Notes**."]

    def img_md(sd: SlideData, k: int, key: str, alt: str, where: str) -> str:
        info = images[key]
        label = f"Slide {sd.number} – {where}image {k}"
        if info.raw_only:
            line = f"[{label} ({info.ext.upper()} file that could not be converted to PNG; see the full-slide snapshot)](images/{info.name})"
        else:
            line = f"![{label}](images/{info.name})"
        extra = []
        others = [n for n in info.slides if n != sd.number]
        if others:
            extra.append(f"Same image also on slide{'s' if len(others) > 1 else ''} {slides_text(others)}.")
        if alt:
            extra.append(f"Alt text: {alt}")
        return line + ("\n*" + " ".join(extra) + "*" if extra else "")

    current_section = None
    for sd in slides:
        if sd.section and sd.section != current_section:
            current_section = sd.section
            out.append(f"# Section: {sd.section}")
        head = f"## Slide {sd.number}" + (" (hidden)" if sd.hidden else "") + (f": {sd.title}" if sd.title else "")
        out += ["---", head]
        if sd.number in snapped:
            why = f" ({', '.join(sd.snapshot_reasons)})" if sd.snapshot_reasons else ""
            out.append(f"![Slide {sd.number} – full slide{why}](images/slide{sd.number}_full.png)")
        k = 0

        def emit(blocks, where=""):
            nonlocal k
            parts = []
            for blk in blocks:
                if blk[0] == "md":
                    parts.append(blk[1])
                elif images[blk[1]].slides:  # tiny spacer images were dropped
                    k += 1
                    alt = blk[2] if blk[2] != "background image" else ""
                    parts.append(img_md(sd, k, blk[1], alt, "background " if blk[2] == "background image" else where))
            return parts

        body = emit(sd.blocks)
        out += body if body else ["_(no text on this slide)_"]
        if sd.offslide:
            off = emit(sd.offslide)
            if off:
                out.append("*Off-slide content (sits outside the visible slide area):*")
                out += off
        out.append("### Speaker Notes")
        notes = emit(sd.notes, "notes ")
        out += notes if notes else ["_(none)_"]

    if used:
        rows = [["File", "Slide(s)", "Original format", "Original size (px)"]]
        for info in sorted(used, key=lambda i: (i.slides[0], i.ordinal)):
            size = f"{info.size[0]}×{info.size[1]}" if info.size else "?"
            rows.append([f"[{info.name}](images/{info.name})", slides_text(info.slides), info.ext, size])
        out += ["---", "## Image Index", md_table(rows)]
    return "\n\n".join(out).rstrip() + "\n"


# --------------------------------------------------------------------------- driver

@dataclass
class Result:
    outdir: Path
    slides: int
    images: int
    notes: int
    snapshots: int
    warnings: list


def convert(src: Path, snapshots: str = "auto") -> Result:
    """snapshots: 'auto' (only slides with pictures/charts/diagrams), 'all', or 'none'."""
    # /tmp, not $TMPDIR: sandboxed PowerPoint can open files under /tmp via AppleScript, but asks the
    # user for a manual "Grant File Access" click for anything under /var/folders.
    workdir = Path(tempfile.mkdtemp(prefix="pptx2md-", dir="/tmp"))
    outdir = None
    try:
        deck_path = prepare_input(src, workdir)
        prs = Presentation(deck_path)
        ex = Extractor(prs)
        slides = ex.extract()
        warnings: list[str] = []

        for sd in slides:
            for key in sd.image_keys():
                if sd.number not in ex.images[key].slides:
                    ex.images[key].slides.append(sd.number)
        for info in ex.images.values():
            if info.slides:
                decode_image(info, workdir)
                if info.size and max(info.size) < MIN_IMAGE_PX:
                    info.slides = []
        counters = Counter()
        for sd in slides:
            for key in sd.image_keys():
                info = ex.images[key]
                if info.slides and not info.ordinal and info.slides[0] == sd.number:
                    counters[sd.number] += 1
                    info.ordinal = counters[sd.number]
                    ext = "png" if not info.raw_only else (info.ext or "bin")
                    info.name = f"slide{slide_list_str(info.slides)}_img{info.ordinal:02d}.{ext}"

        outdir = make_outdir(src)
        img_dir = outdir / "images"
        for info in ex.images.values():
            if not info.slides:
                continue
            if info.raw_only:
                (img_dir / info.name).write_bytes(info.blob)
            else:
                add_label(info.pil, slides_label(info.slides)).save(img_dir / info.name)
                info.pil = None

        for sd in slides:
            sd.snapshot_reasons = snapshot_reasons(sd, ex.images)
        if snapshots == "all":
            wanted = {sd.number for sd in slides}
        elif snapshots == "auto":
            wanted = {sd.number for sd in slides if sd.snapshot_reasons}
        else:
            wanted = set()
        snapped: set[int] = set()
        if wanted:  # text-only decks never touch PowerPoint
            try:
                for no, img in render_snapshots(deck_path, len(slides), any(s.hidden for s in slides), workdir, wanted):
                    add_label(img, f"Slide {no}").save(img_dir / f"slide{no}_full.png")
                    snapped.add(no)
            except Exception as e:
                log.warning("snapshots failed: %s", e)
                warnings.append(f"Full-slide snapshots skipped: {e}")
                snapped.clear()
                for p in img_dir.glob("slide*_full.png"):
                    p.unlink()

        md = build_markdown(src, prs, slides, ex.images, snapped, ex.footers, warnings)
        (outdir / f"{src.stem}.md").write_text(md, encoding="utf-8")
        n_imgs = sum(1 for i in ex.images.values() if i.slides)
        return Result(outdir, len(slides), n_imgs, sum(1 for s in slides if s.notes_plain.strip()), len(snapped), warnings)
    except Exception:
        if outdir is not None:
            shutil.rmtree(outdir, ignore_errors=True)
        raise
    finally:
        shutil.rmtree(workdir, ignore_errors=True)


def osa(script: str):
    subprocess.run(["/usr/bin/osascript", "-e", script], capture_output=True, timeout=120)


def notify(msg: str, enabled: bool):
    log.info(msg)
    if enabled:
        osa(f"display notification {_as(msg)} with title {_as(APP_NAME)}")


def alert(msg: str, enabled: bool):
    log.error(msg)
    if enabled:
        osa(f"display alert {_as(APP_NAME)} message {_as(msg)} as warning giving up after 120")


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("files", nargs="+", type=Path)
    snap = ap.add_mutually_exclusive_group()
    snap.add_argument("--no-snapshots", action="store_true", help="skip full-slide snapshots (no PowerPoint needed)")
    snap.add_argument("--all-snapshots", action="store_true", help="snapshot every slide, including text-only ones")
    ap.add_argument("--no-open", action="store_true", help="don't open the output folder when done")
    ap.add_argument("--no-notify", action="store_true", help="no notifications or alert dialogs")
    args = ap.parse_args(argv)

    LOG_PATH.parent.mkdir(parents=True, exist_ok=True)
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s",
                        handlers=[logging.FileHandler(LOG_PATH), logging.StreamHandler(sys.stderr)])
    show = not args.no_notify
    failures = 0
    for src in args.files:
        src = src.expanduser().resolve()
        if src.is_dir() or src.suffix.lower() not in OOXML_SUFFIXES | LEGACY_SUFFIXES:
            alert(f"Skipped “{src.name}”: not a PowerPoint file (.pptx, .ppt, .pptm, .ppsx, .potx…).", show)
            failures += 1
            continue
        notify(f"Converting {src.name}…", show)
        try:
            r = convert(src, snapshots="none" if args.no_snapshots else "all" if args.all_snapshots else "auto")
        except Exception as e:
            log.error("failed on %s\n%s", src, traceback.format_exc())
            alert(f"Could not convert “{src.name}”:\n{e}\n\nDetails: {LOG_PATH}", show)
            failures += 1
            continue
        msg = f"{src.stem}: {r.slides} slides, {r.images} images, {r.snapshots} slide snapshots, notes on {r.notes} slides"
        if r.warnings:
            msg += " (" + "; ".join(r.warnings) + ")"
        notify(msg, show)
        print(r.outdir)
        if not args.no_open:
            subprocess.run(["/usr/bin/open", str(r.outdir)])
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
